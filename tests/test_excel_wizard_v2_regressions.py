"""Focused regressions for Excel Wizard V2 workbook semantics.

These cases cover behavior that is easy to advertise in the UI while still
getting subtly wrong in a generated workbook: row-based DATE formulas,
standalone MATCH positions, two-axis pivot reports, auditable aggregate output,
and no-store response headers for every workbook endpoint.
"""

from __future__ import annotations

from io import BytesIO
import json
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main
from tests.test_excel_wizard_v2 import _workbook_bytes


class ExcelWizardV2RegressionTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(main.app, base_url="https://testserver")
        self.addCleanup(self.client.close)
        self.client.cookies.set(
            "gauth",
            main._encrypt_token(json.dumps({"email": "dataprovido@gmail.com", "login_type": "email"})),
        )
        response = self.client.post(
            "/api/excel-wizard/upload",
            files={
                "file": (
                    "commerce.xlsx",
                    _workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.workbook_id = response.json()["workbook_id"]
        self.upload_response = response

    def tearDown(self):
        self.client.delete(f"/api/excel-wizard/workbook/{self.workbook_id}")

    def execute(self, operation: dict, command: str = "Uygula") -> dict:
        response = self.client.post(
            "/api/excel-wizard/execute",
            json={
                "workbook_id": self.workbook_id,
                "command": command,
                "mode": "operations",
                "operation": operation,
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload.get("status"), "success", payload)
        return payload

    def download_workbook(self):
        response = self.client.get(f"/api/excel-wizard/download/{self.workbook_id}")
        self.assertEqual(response.status_code, 200, response.text)
        return load_workbook(BytesIO(response.content), data_only=False)

    def test_date_uses_each_rows_year_month_and_day_columns(self):
        payload = self.execute(
            {
                "id": "date",
                "year_column": "Year",
                "month_column": "Month",
                "day_column": "Day",
                "target_column": "Composed Date",
            },
            "Year Month Day sütunlarından tarih oluştur",
        )
        self.assertTrue(payload["result"]["is_mutation"])
        self.assertEqual(payload["result"]["metrics"]["rows_updated"], 5)

        sheet = self.download_workbook()["Commerce Data"]
        headers = [cell.value for cell in sheet[1]]
        target = headers.index("Composed Date") + 1
        year = headers.index("Year") + 1
        month = headers.index("Month") + 1
        day = headers.index("Day") + 1
        for row in range(2, 7):
            expected = (
                f"=DATE({sheet.cell(row, year).coordinate},"
                f"{sheet.cell(row, month).coordinate},{sheet.cell(row, day).coordinate})"
            )
            self.assertEqual(sheet.cell(row, target).value, expected)

    def test_standalone_match_returns_one_based_data_position_without_lookup_sheet(self):
        payload = self.execute(
            {"id": "index_match", "source_column": "Brand", "match_value": "Samsung"},
            "Brand sütununda Samsung kaçıncı sırada?",
        )
        result = payload["result"]
        self.assertEqual(result["metrics"]["match_position"], 3)
        self.assertFalse(result["is_mutation"])
        self.assertEqual(result["columns_used"], ["Brand"])

    def test_pivot_supports_rows_columns_and_values_axes(self):
        payload = self.execute(
            {
                "id": "pivot",
                "index": "Brand",
                "columns": "Category",
                "values": "Revenue",
                "aggfunc": "sum",
                "sheet_name": "Brand_Category_Pivot",
            },
            "Brand satır, Category sütun ve Revenue değer pivotu oluştur",
        )
        self.assertEqual(payload["result"]["columns_used"], ["Brand", "Category", "Revenue"])

        pivot = self.download_workbook()["Brand_Category_Pivot"]
        headers = [cell.value for cell in pivot[1]]
        rows = {
            pivot.cell(row, 1).value: {
                headers[column - 1]: pivot.cell(row, column).value
                for column in range(2, pivot.max_column + 1)
            }
            for row in range(2, pivot.max_row + 1)
        }
        self.assertEqual(rows["Apple"]["Telefon"], 2400)
        self.assertEqual(rows["Apple"]["Computers"], 2000)
        self.assertEqual(rows["Samsung"]["Telefon"], 1500)
        self.assertEqual(rows["Samsung"]["Home"], 600)

    def test_aggregate_result_and_excel_formula_are_written_to_download(self):
        payload = self.execute(
            {"id": "sum", "column": "Revenue"},
            "Revenue sütununu topla",
        )
        self.assertTrue(payload["result"]["is_mutation"])
        self.assertEqual(payload["result"]["metrics"]["sum"], 7000)

        workbook = self.download_workbook()
        self.assertIn("DataProvido_Results", workbook.sheetnames)
        analysis = workbook["DataProvido_Results"]
        self.assertEqual(analysis["B2"].value, "SUM")
        self.assertEqual(analysis["C2"].value, "Revenue")
        self.assertEqual(analysis["D2"].value, 7000)
        self.assertEqual(analysis["E2"].value, "=SUM('Commerce Data'!$D$2:$D$6)")

    def test_every_workbook_route_sets_private_no_store(self):
        self.assertIn("no-store", self.upload_response.headers.get("cache-control", ""))

        state = self.client.get("/api/excel-wizard/state", params={"workbook_id": self.workbook_id})
        execute = self.client.post(
            "/api/excel-wizard/execute",
            json={
                "workbook_id": self.workbook_id,
                "command": "Revenue toplamı",
                "mode": "operations",
                "operation": {"id": "sum", "column": "Revenue"},
            },
        )
        download = self.client.get(f"/api/excel-wizard/download/{self.workbook_id}")
        delete = self.client.delete(f"/api/excel-wizard/workbook/{self.workbook_id}")
        for response in (state, execute, download, delete):
            with self.subTest(path=response.request.url.path):
                self.assertIn("private", response.headers.get("cache-control", ""))
                self.assertIn("no-store", response.headers.get("cache-control", ""))


if __name__ == "__main__":
    unittest.main()
