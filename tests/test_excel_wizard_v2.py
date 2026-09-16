"""Acceptance tests for the account-scoped, ephemeral Excel Wizard API.

The matrix below is intentionally written in the language of the product
requirement.  A capability advertised by the upload endpoint is a public
promise: the UI may offer it and callers may submit the same capability id in
``operation`` or ``operations``.  The artifact tests complement that breadth
check by opening the downloaded workbook and verifying representative formula,
formatting, validation, pivot-report and chart mutations.

These tests never use a real login, AI provider, network request or persistent
customer file.  Workbooks are built in memory and authenticated with the same
signed console cookie used by the rest of the application test suite.
"""

from __future__ import annotations

from io import BytesIO
import html
import json
import os
from pathlib import Path
import re
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pandas as pd

import main


# One entry for each numbered item in the customer's requirement.  Slash
# groups remain one product requirement but may expose multiple operation ids.
REQUESTED_ACCEPTANCE_MATRIX = {
    1: ("sum",),
    2: ("average",),
    3: ("if",),
    4: ("ifs",),
    5: ("sumproduct",),
    6: ("sumif",),
    7: ("sumifs",),
    8: ("countif",),
    9: ("countifs",),
    10: ("counta",),
    11: ("vlookup",),
    12: ("hlookup",),
    13: ("index_match",),
    14: ("index_match",),
    15: ("lookup",),
    16: ("xlookup",),
    17: ("concat", "textjoin"),
    18: ("left", "right", "mid"),
    19: ("upper", "lower"),
    20: ("substitute", "find"),
    21: ("len",),
    22: ("today", "now"),
    23: ("date",),
    24: ("day", "month", "year"),
    25: ("conditional_format",),
    26: ("filter", "sort"),
    27: ("pivot",),
    28: ("data_validation",),
    29: ("text_to_columns",),
    30: ("chart",),
}

REQUIRED_CAPABILITIES = {
    capability
    for requirement in REQUESTED_ACCEPTANCE_MATRIX.values()
    for capability in requirement
} | {"commerce_analysis"}

# The gallery has exactly 30 product-level buttons.  Some buttons intentionally
# compose several primitive workbook operations, so the hint-to-plan mapping is
# part of the public UI/API contract rather than an implementation detail.
GALLERY_EXPECTED_PLANS = {
    "sum": ("sum",),
    "average": ("average",),
    "if": ("if",),
    "ifs": ("ifs",),
    "sumproduct": ("sumproduct",),
    "sumif": ("sumif",),
    "sumifs": ("sumifs",),
    "countif": ("countif",),
    "countifs": ("countifs",),
    "counta": ("counta",),
    "vlookup": ("vlookup",),
    "hlookup": ("hlookup",),
    "index": ("index_match",),
    "match": ("index_match",),
    "lookup": ("lookup",),
    "xlookup": ("xlookup",),
    "concat": ("concat",),
    "text-slice": ("left", "right", "mid"),
    "case": ("upper", "lower"),
    "substitute-find": ("substitute", "find"),
    "len": ("len",),
    "today-now": ("today", "now"),
    "date": ("date",),
    "date-parts": ("day", "month", "year"),
    "conditional-formatting": ("conditional_format",),
    "filter-sort": ("filter", "sort"),
    "pivot": ("pivot",),
    "validation": ("data_validation",),
    "text-to-columns": ("text_to_columns",),
    "chart": ("chart",),
}


def _workbook_bytes(*, ambiguous_revenue: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commerce Data"
    headers = [
        "Product",
        "Brand",
        "Category",
        "Revenue",
        "Units",
        "Price",
        "Status",
        "OrderDate",
        "Code",
        "MonthKey",
        "SKU",
        "Stock",
        "Quantity",
        "Year",
        "Month",
        "Day",
        "Order Date",
        "Product Info",
        "Order ID",
        "PDP Views",
        "Add to Cart",
        "Transactions",
        "Cost",
        "Margin",
    ]
    if ambiguous_revenue:
        headers.append("Net Revenue")
    sheet.append(headers)
    rows = [
        ["iPhone", "Apple", "Telefon", 2400, 2, 1200, "Active", "2026-01-15", "TR-IST-001", "Jan", "SKU-APL-001", 5, 2, 2026, 1, 15, "2026-01-15", "iPhone-Apple-Pro", "ORD-001", 100, 20, 2, 900, 300],
        ["MacBook", "Apple", "Computers", 2000, 1, 2000, "Active", "2026-02-20", "TR-IST-002", "Feb", "SKU-APL-002", 20, 1, 2026, 2, 20, "2026-02-20", "MacBook-Apple-Air", "ORD-002", 80, 16, 1, 1600, 400],
        ["Galaxy", "Samsung", "Telefon", 1500, 3, 500, "Active", "2026-02-28", "TR-ANK-003", "Mar", "SKU-SAM-003", 7, 3, 2026, 2, 28, "2026-02-28", "Galaxy-Samsung-S24", "ORD-003", 120, 30, 3, 350, 150],
        ["Television", "Samsung", "Home", 600, 2, 300, "Passive", "2026-03-03", "TR-ANK-004", "Jan", "SKU-SAM-004", 15, 2, 2026, 3, 3, "2026-03-03", "Television-Samsung-OLED", "ORD-004", 50, 10, 2, 230, 70],
        ["Watch", "Garmin", "Accessories", 500, 5, 100, "Pending", "2026-03-12", "TR-IZM-005", "Feb", "SKU-GAR-005", 4, 5, 2026, 3, 12, "2026-03-12", "Watch-Garmin-Sport", "ORD-005", 70, 14, 5, 70, 30],
    ]
    if ambiguous_revenue:
        for index, row in enumerate(rows, start=1):
            row.append(row[3] - (index * 25))
    for row in rows:
        sheet.append(row)

    # Preservation sentinel: the V2 engine must modify a workbook rather than
    # flattening it through pandas and losing workbook structure or styling.
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}6"

    lookup = workbook.create_sheet("Lookup")
    lookup.append(["Brand", "Discount"])
    lookup.append(["Apple", 0.05])
    lookup.append(["Samsung", 0.10])
    lookup.append(["Garmin", 0.08])
    lookup["D1"] = "Jan"
    lookup["E1"] = "Feb"
    lookup["F1"] = "Mar"
    lookup["D2"] = 100
    lookup["E2"] = 125
    lookup["F2"] = 150

    horizontal = workbook.create_sheet("Monthly")
    horizontal.append(["Jan", "Feb", "Mar"])
    horizontal.append([100, 125, 150])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _capability_ids(payload: dict) -> set[str]:
    result = set()
    for item in payload.get("capabilities") or []:
        if isinstance(item, str):
            result.add(item)
        elif isinstance(item, dict) and item.get("id"):
            result.add(str(item["id"]))
    return result


def _all_scalar_text(value) -> str:
    if isinstance(value, dict):
        return " ".join(_all_scalar_text(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return " ".join(_all_scalar_text(item) for item in value)
    return "" if value is None else str(value)


def _gallery_buttons() -> list[tuple[str, str]]:
    """Return the exact operation hint and prompt visible in the UI gallery."""
    root = Path(main.__file__).resolve().parent
    template = (root / "templates" / "excel_wizard.html").read_text(encoding="utf-8")
    return [
        (operation, html.unescape(prompt))
        for operation, prompt in re.findall(
            r'<button\b[^>]*data-xw-operation="([^"]+)"[^>]*data-xw-prompt="([^"]*)"',
            template,
        )
    ]


class ExcelWizardV2ContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.routes = {getattr(route, "path", None) for route in main.app.routes}
        expected = {
            "/api/excel-wizard/upload",
            "/api/excel-wizard/state",
            "/api/excel-wizard/execute",
            "/api/excel-wizard/workbook/{workbook_id}",
            "/api/excel-wizard/download/{workbook_id}",
        }
        missing = expected - cls.routes
        if missing:
            raise AssertionError(f"Excel Wizard V2 routes are not integrated: {sorted(missing)}")

    def setUp(self):
        self.client = TestClient(main.app, base_url="https://testserver")
        self.addCleanup(self.client.close)
        self.created: list[tuple[str, str]] = []
        self.sign_in("dataprovido@gmail.com")

    def tearDown(self):
        # Delete every workbook through its public owner-bound API.  This also
        # ensures the suite itself never leaves uploaded customer data behind.
        for owner, workbook_id in reversed(self.created):
            self.sign_in(owner)
            self.client.delete(f"/api/excel-wizard/workbook/{workbook_id}")

    def sign_in(self, email: str) -> None:
        self.client.cookies.set(
            "gauth",
            main._encrypt_token(json.dumps({"email": email, "login_type": "email"})),
        )

    def upload(self, *, ambiguous_revenue: bool = False, filename: str = "commerce.xlsx") -> dict:
        response = self.client.post(
            "/api/excel-wizard/upload",
            files={
                "file": (
                    filename,
                    _workbook_bytes(ambiguous_revenue=ambiguous_revenue),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload.get("status"), "success", payload)
        self.created.append(("dataprovido@gmail.com", payload["workbook_id"]))
        return payload

    def execute(self, workbook_id: str, **payload) -> dict:
        response = self.client.post(
            "/api/excel-wizard/execute",
            json={"workbook_id": workbook_id, **payload},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_all_v2_routes_require_a_signed_console_user(self):
        self.client.cookies.clear()
        calls = [
            self.client.post(
                "/api/excel-wizard/upload",
                files={"file": ("commerce.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            ),
            self.client.get("/api/excel-wizard/state", params={"workbook_id": "unknown"}),
            self.client.post(
                "/api/excel-wizard/execute",
                json={"workbook_id": "unknown", "command": "Revenue toplamı", "mode": "operations"},
            ),
            self.client.get("/api/excel-wizard/download/unknown"),
            self.client.delete("/api/excel-wizard/workbook/unknown"),
        ]
        self.assertTrue(all(response.status_code == 401 for response in calls), [response.text for response in calls])

    def test_upload_returns_opaque_id_profile_preview_and_complete_capability_matrix(self):
        payload = self.upload()
        workbook_id = payload["workbook_id"]
        self.assertRegex(workbook_id, r"^[A-Za-z0-9_-]{16,}$")
        self.assertNotIn("/", workbook_id)
        self.assertEqual(payload["filename"], "commerce.xlsx")
        self.assertIsInstance(payload.get("profile"), dict)
        self.assertIsInstance(payload.get("preview"), dict)
        self.assertIn("Revenue", payload["preview"].get("columns", []))
        self.assertEqual(payload["preview"].get("total_rows"), 5)
        self.assertEqual(len(REQUESTED_ACCEPTANCE_MATRIX), 30)
        self.assertTrue(
            REQUIRED_CAPABILITIES.issubset(_capability_ids(payload)),
            sorted(REQUIRED_CAPABILITIES - _capability_ids(payload)),
        )
        # Filesystem details must never cross the API boundary.
        self.assertNotRegex(_all_scalar_text(payload), r"(?:data/uploads|/tmp/|\\tmp\\)")

    def test_all_visible_gallery_prompts_map_to_the_expected_primitive_operations(self):
        """The label a customer clicks must be the plan the engine executes.

        This is deliberately independent from keyword inference: ``operation_hint``
        is the stable bridge between the 30 visible product buttons and the
        primitive workbook operations.  Composite gallery items must expand to
        every advertised operation, in order.
        """
        from functions.excel_wizard import build_operation_from_command

        buttons = _gallery_buttons()
        self.assertEqual(len(buttons), 30, buttons)
        self.assertEqual({operation for operation, _ in buttons}, set(GALLERY_EXPECTED_PLANS))

        workbook_id = self.upload()["workbook_id"]
        record = main.excel_workbook_store.get("dataprovido@gmail.com", workbook_id)
        for operation_hint, prompt in buttons:
            with self.subTest(operation_hint=operation_hint, prompt=prompt):
                plan = build_operation_from_command(
                    prompt,
                    "operations",
                    record,
                    operation_hint=operation_hint,
                )
                operations = plan.get("_batch") or [plan]
                actual = tuple(str(operation.get("id") or operation.get("type")) for operation in operations)
                self.assertEqual(actual, GALLERY_EXPECTED_PLANS[operation_hint], plan)

    def test_all_visible_gallery_prompts_execute_or_ask_an_actionable_question(self):
        """Exercise the browser's exact prompt/hint payload through FastAPI."""
        buttons = _gallery_buttons()
        self.assertEqual(len(buttons), 30, buttons)

        for operation_hint, prompt in buttons:
            with self.subTest(operation_hint=operation_hint, prompt=prompt):
                workbook_id = self.upload()["workbook_id"]
                response = self.client.post(
                    "/api/excel-wizard/execute",
                    json={
                        "workbook_id": workbook_id,
                        "command": prompt,
                        "mode": "operations",
                        "operation_hint": operation_hint,
                    },
                )
                self.assertEqual(response.status_code, 200, response.text)
                payload = response.json()
                self.assertIn(payload.get("status"), {"success", "needs_clarification"}, payload)

                expected = GALLERY_EXPECTED_PLANS[operation_hint]
                if payload["status"] == "needs_clarification":
                    self.assertTrue(payload.get("clarification_id"), payload)
                    self.assertTrue(payload.get("role"), payload)
                    self.assertTrue(payload.get("question"), payload)
                    self.assertTrue(payload.get("options"), payload)
                    continue

                result = payload["result"]
                if len(expected) == 1:
                    self.assertEqual(result.get("operation_id"), expected[0], result)
                else:
                    self.assertEqual(result.get("operation_id"), "batch", result)
                    step_text = "\n".join(result.get("steps") or [])
                    for operation_id in expected:
                        self.assertIn(f"İşlem planı: {operation_id}", step_text)
                self.assertTrue(payload.get("download_url"), payload)

    def test_gallery_aggregate_prompts_use_the_named_columns_and_conditions(self):
        expected_metrics = {
            "sum": ("sum", 7000),
            "average": ("average", 820),
            "sumproduct": ("sumproduct", 7000),
            "sumif": ("sumif", 4400),
            "sumifs": ("sumifs", 2400),
            "countif": ("countif", 3),
            "countifs": ("countifs", 1),
            "counta": ("counta", 5),
        }
        prompts = dict(_gallery_buttons())
        for operation_hint, (metric_name, expected_value) in expected_metrics.items():
            with self.subTest(operation_hint=operation_hint):
                workbook_id = self.upload()["workbook_id"]
                payload = self.execute(
                    workbook_id,
                    command=prompts[operation_hint],
                    mode="operations",
                    operation_hint=operation_hint,
                )
                self.assertEqual(payload.get("status"), "success", payload)
                self.assertAlmostEqual(float(payload["result"]["metrics"][metric_name]), expected_value)

    def test_state_is_owner_bound_and_does_not_leak_workbook_existence(self):
        payload = self.upload()
        workbook_id = payload["workbook_id"]

        self.sign_in("myasamkaradag@gmail.com")
        for response in (
            self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}),
            self.client.post(
                "/api/excel-wizard/execute",
                json={"workbook_id": workbook_id, "command": "Revenue toplamı", "mode": "operations"},
            ),
            self.client.get(f"/api/excel-wizard/download/{workbook_id}"),
            self.client.delete(f"/api/excel-wizard/workbook/{workbook_id}"),
        ):
            self.assertEqual(response.status_code, 404, response.text)

        self.sign_in("dataprovido@gmail.com")
        self.assertEqual(
            self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).status_code,
            200,
        )

    def test_delete_is_immediate_and_removes_state_execute_and_download_access(self):
        payload = self.upload()
        workbook_id = payload["workbook_id"]
        response = self.client.delete(f"/api/excel-wizard/workbook/{workbook_id}")
        self.assertEqual(response.status_code, 200, response.text)
        self.created.clear()
        self.assertEqual(
            self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                "/api/excel-wizard/execute",
                json={"workbook_id": workbook_id, "command": "Revenue toplamı", "mode": "operations"},
            ).status_code,
            404,
        )
        self.assertEqual(self.client.get(f"/api/excel-wizard/download/{workbook_id}").status_code, 404)

    def test_expired_workbook_is_cleaned_without_sleeping_or_disk_fallback(self):
        payload = self.upload()
        workbook_id = payload["workbook_id"]
        store = main.excel_workbook_store
        count_before = store.count()
        self.assertGreaterEqual(count_before, 1)
        self.assertEqual(store.expire(workbook_id), 1)
        self.assertEqual(store.cleanup(), 1)
        self.assertEqual(store.count(), count_before - 1)
        self.created.clear()

        self.assertEqual(
            self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).status_code,
            404,
        )
        self.assertEqual(self.client.get(f"/api/excel-wizard/download/{workbook_id}").status_code, 404)

    def test_upload_is_memory_only_and_does_not_touch_legacy_active_files(self):
        upload_dir = Path(main.__file__).resolve().parent / "data" / "uploads"
        legacy_names = {
            "active_uploaded_dataset.xlsx",
            "active_original_dataset.xlsx",
            "active_meta.json",
        }

        def snapshot():
            return {
                name: (upload_dir / name).read_bytes() if (upload_dir / name).exists() else None
                for name in legacy_names
            }

        before = snapshot()
        payload = self.upload()
        workbook_id = payload["workbook_id"]
        self.execute(
            workbook_id,
            command="Revenue sütununu topla",
            mode="operations",
            operation={"id": "sum", "column": "Revenue"},
        )
        self.client.get(f"/api/excel-wizard/download/{workbook_id}")
        self.assertEqual(snapshot(), before)

    def test_csv_is_supported_but_invalid_and_unsupported_files_are_rejected(self):
        csv_response = self.client.post(
            "/api/excel-wizard/upload",
            files={"file": ("small.csv", b"Brand,Revenue\nApple,100\nSamsung,80\n", "text/csv")},
        )
        self.assertEqual(csv_response.status_code, 200, csv_response.text)
        csv_payload = csv_response.json()
        self.created.append(("dataprovido@gmail.com", csv_payload["workbook_id"]))
        self.assertEqual(csv_payload["preview"]["total_rows"], 2)

        for name, contents in (("notes.txt", b"not a workbook"), ("broken.xlsx", b"not-a-zip")):
            with self.subTest(filename=name):
                response = self.client.post(
                    "/api/excel-wizard/upload",
                    files={"file": (name, contents, "application/octet-stream")},
                )
                self.assertIn(response.status_code, {400, 415}, response.text)

    def test_execute_success_has_stable_result_contract(self):
        workbook_id = self.upload()["workbook_id"]
        payload = self.execute(
            workbook_id,
            command="Revenue sütununun toplamını hesapla",
            mode="operations",
            operation={"id": "sum", "column": "Revenue"},
        )
        self.assertEqual(payload.get("status"), "success", payload)
        result = payload.get("result") or {}
        for key in ("summary", "steps", "metrics", "columns_used", "preview_rows", "is_mutation"):
            self.assertIn(key, result)
        self.assertIn("Revenue", result["columns_used"])
        self.assertRegex(_all_scalar_text(result), r"(?:7[.,]?000|7000)")
        self.assertIsInstance(payload.get("preview"), dict)
        self.assertEqual(payload.get("download_url"), f"/api/excel-wizard/download/{workbook_id}")

    def test_operations_mode_mutates_same_workbook_and_download_is_real_xlsx(self):
        workbook_id = self.upload()["workbook_id"]
        payload = self.execute(
            workbook_id,
            command="Brand sütununu büyük harfe çevirip Brand_Upper sütununa yaz",
            mode="operations",
            operation={
                "id": "upper",
                "source_column": "Brand",
                "target_column": "Brand_Upper",
            },
        )
        self.assertEqual(payload.get("status"), "success", payload)
        self.assertTrue(payload["result"]["is_mutation"])

        state = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id})
        self.assertEqual(state.status_code, 200, state.text)
        self.assertTrue(state.json()["is_modified"])
        self.assertIn("Brand_Upper", state.json()["preview"]["columns"])

        download = self.client.get(f"/api/excel-wizard/download/{workbook_id}")
        self.assertEqual(download.status_code, 200, download.text)
        self.assertEqual(
            download.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertRegex(download.headers.get("content-disposition", ""), r"DataProvido_commerce_.*\.xlsx")
        workbook = load_workbook(BytesIO(download.content), data_only=False)
        self.assertIn("Commerce Data", workbook.sheetnames)
        self.assertIn("Lookup", workbook.sheetnames)
        sheet = workbook["Commerce Data"]
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "1F4E78")
        self.assertEqual(sheet.freeze_panes, "A2")
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Brand_Upper", headers)
        target_letter = sheet.cell(row=1, column=headers.index("Brand_Upper") + 1).column_letter
        values = [sheet[f"{target_letter}{row}"].value for row in range(2, 7)]
        self.assertTrue(all(str(value).upper().startswith(("=UPPER", "APPLE", "SAMSUNG", "GARMIN")) for value in values))

    def test_all_conditional_aggregate_families_calculate_exact_values(self):
        workbook_id = self.upload()["workbook_id"]
        operations = [
            {"id": "sum", "column": "Revenue"},
            {"id": "average", "column": "Revenue"},
            {"id": "sumproduct", "columns": ["Units", "Price"]},
            {
                "id": "sumif",
                "sum_column": "Revenue",
                "criteria": [{"column": "Brand", "operator": "equals", "value": "Apple"}],
            },
            {
                "id": "sumifs",
                "sum_column": "Revenue",
                "criteria": [
                    {"column": "Brand", "operator": "equals", "value": "Apple"},
                    {"column": "Category", "operator": "equals", "value": "Telefon"},
                ],
            },
            {
                "id": "countif",
                "criteria": [{"column": "Status", "operator": "equals", "value": "Active"}],
            },
            {
                "id": "countifs",
                "criteria": [
                    {"column": "Brand", "operator": "equals", "value": "Samsung"},
                    {"column": "Status", "operator": "equals", "value": "Active"},
                ],
            },
            {"id": "counta", "column": "Product"},
        ]
        payload = self.execute(
            workbook_id,
            command="Toplam, ortalama ve koşullu hesaplamaları uygula",
            mode="operations",
            operations=operations,
        )
        self.assertEqual(payload.get("status"), "success", payload)
        metrics = payload["result"]["metrics"]
        expected = {
            "sum_sum": 7000,
            "average_average": 1400,
            "sumproduct_sumproduct": 7000,
            "sumif_sumif": 4400,
            "sumifs_sumifs": 2400,
            "countif_countif": 3,
            "countifs_countifs": 1,
            "counta_counta": 5,
        }
        for key, value in expected.items():
            with self.subTest(metric=key):
                self.assertAlmostEqual(float(metrics[key]), value)
        # Operations-mode aggregate results are persisted in the downloadable
        # workbook's DataProvido_Results audit sheet.
        self.assertTrue(payload["result"]["is_mutation"])

    def test_formula_text_and_date_families_write_auditable_excel_formulas(self):
        workbook_id = self.upload()["workbook_id"]
        operations = [
            {
                "id": "if",
                "source_column": "Revenue",
                "target_column": "Revenue_Band",
                "operator": ">=",
                "value": 1500,
                "true_value": "High",
                "false_value": "Standard",
            },
            {
                "id": "ifs",
                "source_column": "Revenue",
                "target_column": "Revenue_Tier",
                "cases": [
                    {"operator": ">", "value": 2000, "result": "High"},
                    {"operator": ">", "value": 1000, "result": "Medium"},
                    {"operator": ">=", "value": 0, "result": "Low"},
                ],
            },
            {"id": "concat", "columns": ["Product", "Brand"], "delimiter": " - ", "target_column": "Product_Brand"},
            {"id": "textjoin", "columns": ["Brand", "Category"], "delimiter": " / ", "target_column": "Brand_Category"},
            {"id": "left", "source_column": "Code", "count": 2, "target_column": "Country"},
            {"id": "right", "source_column": "Code", "count": 3, "target_column": "Sequence"},
            {"id": "mid", "source_column": "Code", "start": 4, "count": 3, "target_column": "City"},
            {"id": "upper", "source_column": "Brand", "target_column": "Brand_Upper"},
            {"id": "lower", "source_column": "Brand", "target_column": "Brand_Lower"},
            {"id": "substitute", "source_column": "Code", "old": "-", "new": "/", "target_column": "Code_Slashed"},
            {"id": "find", "source_column": "Code", "text": "-", "target_column": "First_Dash"},
            {"id": "len", "source_column": "Product", "target_column": "Product_Length"},
            {"id": "today", "target_column": "Report_Date"},
            {"id": "now", "target_column": "Report_Timestamp"},
            {"id": "date", "year": 2026, "month": 9, "day": 16, "target_column": "Fixed_Date"},
            {"id": "day", "source_column": "OrderDate", "target_column": "Order_Day"},
            {"id": "month", "source_column": "OrderDate", "target_column": "Order_Month"},
            {"id": "year", "source_column": "OrderDate", "target_column": "Order_Year"},
        ]
        payload = self.execute(
            workbook_id,
            command="Formül, metin ve tarih kolonlarını oluştur",
            mode="operations",
            operations=operations,
        )
        self.assertEqual(payload.get("status"), "success", payload)
        self.assertTrue(payload["result"]["is_mutation"])

        workbook = load_workbook(
            BytesIO(self.client.get(f"/api/excel-wizard/download/{workbook_id}").content),
            data_only=False,
        )
        sheet = workbook["Commerce Data"]
        header_to_column = {cell.value: cell.column_letter for cell in sheet[1]}
        expected_formulas = {
            "Revenue_Band": "=IF(",
            "Revenue_Tier": "=IFS(",
            "Product_Brand": "=CONCAT(",
            "Brand_Category": "=TEXTJOIN(",
            "Country": "=LEFT(",
            "Sequence": "=RIGHT(",
            "City": "=MID(",
            "Brand_Upper": "=UPPER(",
            "Brand_Lower": "=LOWER(",
            "Code_Slashed": "=SUBSTITUTE(",
            "First_Dash": "=IFERROR(FIND(",
            "Product_Length": "=LEN(",
            "Report_Date": "=TODAY(",
            "Report_Timestamp": "=NOW(",
            "Fixed_Date": "=DATE(",
            "Order_Day": "=DAY(",
            "Order_Month": "=MONTH(",
            "Order_Year": "=YEAR(",
        }
        for header, prefix in expected_formulas.items():
            with self.subTest(header=header):
                self.assertIn(header, header_to_column)
                self.assertTrue(str(sheet[f"{header_to_column[header]}2"].value).upper().startswith(prefix), sheet[f"{header_to_column[header]}2"].value)

    def test_lookup_families_preserve_lookup_sheet_and_write_expected_formulas(self):
        workbook_id = self.upload()["workbook_id"]
        operations = [
            {
                "id": operation_id,
                "source_column": "Brand",
                "lookup_sheet": "Lookup",
                "lookup_key_column": "Brand",
                "return_column": "Discount",
                "target_column": f"Discount_{operation_id}",
            }
            for operation_id in ("vlookup", "index_match", "lookup", "xlookup")
        ]
        operations.append(
            {
                "id": "hlookup",
                "source_column": "MonthKey",
                "lookup_sheet": "Monthly",
                "return_row": 2,
                "target_column": "Monthly_Target",
            }
        )
        payload = self.execute(
            workbook_id,
            command="Brand değerlerini Lookup sayfasından eşleştir",
            mode="operations",
            operations=operations,
        )
        self.assertEqual(payload.get("status"), "success", payload)

        workbook = load_workbook(
            BytesIO(self.client.get(f"/api/excel-wizard/download/{workbook_id}").content),
            data_only=False,
        )
        self.assertIn("Lookup", workbook.sheetnames)
        sheet = workbook["Commerce Data"]
        header_to_column = {cell.value: cell.column_letter for cell in sheet[1]}
        expected = {
            "Discount_vlookup": "VLOOKUP(",
            "Discount_index_match": "INDEX(",
            "Discount_lookup": "LOOKUP(",
            "Discount_xlookup": "XLOOKUP(",
            "Monthly_Target": "HLOOKUP(",
        }
        for header, token in expected.items():
            with self.subTest(header=header):
                formula = str(sheet[f"{header_to_column[header]}2"].value).upper()
                self.assertIn(token, formula)
        preview_row = payload["preview"]["rows"][0]
        self.assertEqual(preview_row["Discount_vlookup"], 0.05)
        self.assertEqual(preview_row["Discount_index_match"], 0.05)
        self.assertEqual(preview_row["Discount_lookup"], 0.05)
        self.assertEqual(preview_row["Discount_xlookup"], 0.05)
        self.assertEqual(preview_row["Monthly_Target"], 100)
        self.assertEqual(payload["preview"]["rows"][1]["Monthly_Target"], 125)
        self.assertEqual(payload["preview"]["rows"][2]["Monthly_Target"], 150)

    def test_text_to_columns_creates_named_columns_without_losing_source(self):
        workbook_id = self.upload()["workbook_id"]
        payload = self.execute(
            workbook_id,
            command="Code sütununu üç sütuna ayır",
            mode="operations",
            operation={
                "id": "text_to_columns",
                "column": "Code",
                "delimiter": "-",
                "target_columns": ["Country", "City", "Sequence"],
            },
        )
        self.assertEqual(payload.get("status"), "success", payload)
        self.assertEqual(payload["preview"]["rows"][0]["Code"], "TR-IST-001")
        self.assertEqual(payload["preview"]["rows"][0]["Country"], "TR")
        self.assertEqual(payload["preview"]["rows"][0]["City"], "IST")
        self.assertEqual(payload["preview"]["rows"][0]["Sequence"], "001")

    def test_representative_excel_artifacts_are_preserved_in_download(self):
        workbook_id = self.upload()["workbook_id"]
        operations = [
            {
                "id": "conditional_format",
                "column": "Revenue",
                "operator": "greaterThan",
                "value": 1500,
                "color": "FDE68A",
            },
            {"id": "data_validation", "column": "Status", "values": ["Active", "Paused"]},
            {
                "id": "pivot",
                "index": "Brand",
                "values": ["Revenue"],
                "aggfunc": "sum",
                "sheet_name": "Brand_Revenue",
            },
            {
                "id": "chart",
                "category_column": "Brand",
                "value_columns": ["Revenue"],
                "chart_type": "bar",
                "title": "Revenue by brand",
            },
        ]
        payload = self.execute(
            workbook_id,
            command="Revenue için koşullu biçimlendirme, doğrulama, pivot ve grafik oluştur",
            mode="operations",
            operations=operations,
        )
        self.assertEqual(payload.get("status"), "success", payload)

        download = self.client.get(f"/api/excel-wizard/download/{workbook_id}")
        workbook = load_workbook(BytesIO(download.content), data_only=False)
        sheet = workbook["Commerce Data"]
        self.assertGreater(len(sheet.conditional_formatting), 0)
        self.assertGreater(sheet.data_validations.count, 0)
        self.assertGreater(len(sheet._charts), 0)
        chart_data_sheets = [name for name in workbook.sheetnames if name.startswith("Chart_Analysis")]
        self.assertEqual(len(chart_data_sheets), 1)
        chart_data = workbook[chart_data_sheets[0]]
        chart_rows = {chart_data.cell(row, 1).value: chart_data.cell(row, 2).value for row in range(2, chart_data.max_row + 1)}
        self.assertEqual(chart_rows["Apple"], 4400)
        self.assertEqual(chart_rows["Samsung"], 2100)
        self.assertIn("Brand_Revenue", workbook.sheetnames)
        pivot = workbook["Brand_Revenue"]
        pivot_text = " ".join(str(cell.value) for row in pivot.iter_rows() for cell in row if cell.value is not None)
        self.assertIn("Apple", pivot_text)
        self.assertRegex(pivot_text, r"(?:4[.,]?400|4400)")

    def test_filter_and_sort_can_be_combined_deterministically(self):
        workbook_id = self.upload()["workbook_id"]
        payload = self.execute(
            workbook_id,
            command="Active satırları filtrele ve Revenue değerine göre azalan sırala",
            mode="operations",
            operations=[
                {
                    "id": "filter",
                    "criteria": [{"column": "Status", "operator": "equals", "value": "Active"}],
                },
                {"id": "sort", "column": "Revenue", "ascending": False},
            ],
        )
        self.assertEqual(payload.get("status"), "success", payload)
        rows = payload["preview"]["rows"]
        self.assertEqual([row["Revenue"] for row in rows], [2400, 2000, 1500])
        self.assertTrue(all(row["Status"] == "Active" for row in rows))

    def test_commerce_mode_answers_ecommerce_question_without_mutating_workbook(self):
        workbook_id = self.upload()["workbook_id"]
        before = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        payload = self.execute(
            workbook_id,
            command="Revenue'a göre en performanslı marka hangisi?",
            mode="commerce",
            operation={
                "id": "commerce_analysis",
                "metric_column": "Revenue",
                "group_column": "Brand",
                "top_n": 5,
            },
        )
        self.assertEqual(payload.get("status"), "success", payload)
        result = payload["result"]
        self.assertFalse(result["is_mutation"])
        self.assertTrue({"Revenue", "Brand"}.issubset(set(result["columns_used"])))
        self.assertIn("Apple", _all_scalar_text(result))
        self.assertRegex(_all_scalar_text(result), r"(?:4[.,]?400|4400)")

        after = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        self.assertFalse(after["is_modified"])
        self.assertEqual(after["preview"], before["preview"])

    def test_commerce_role_matching_covers_standard_ecommerce_funnel_columns(self):
        from functions.excel_wizard import resolve_column

        workbook_id = self.upload()["workbook_id"]
        record = main.excel_workbook_store.get("dataprovido@gmail.com", workbook_id)
        expected = {
            "brand": "Brand",
            "product": "Product",
            "category": "Category",
            "revenue": "Revenue",
            "order": "Order ID",
            "units": "Units",
            "views": "PDP Views",
            "add_to_cart": "Add to Cart",
            "transactions": "Transactions",
            "stock": "Stock",
            "price": "Price",
            "cost": "Cost",
            "margin": "Margin",
        }
        numeric_roles = {"revenue", "units", "views", "add_to_cart", "transactions", "stock", "price", "cost", "margin"}
        for role, expected_column in expected.items():
            with self.subTest(role=role):
                resolved, candidates = resolve_column(record, role, numeric=role in numeric_roles)
                self.assertEqual(resolved, expected_column, (role, resolved, candidates))
                self.assertEqual(candidates, [])

    def test_commerce_role_matching_understands_common_export_heading_variants(self):
        from functions.excel_wizard import resolve_column

        frame = pd.DataFrame({
            "Manufacturer": ["Apple"],
            "Item Name": ["Phone"],
            "Item Category": ["Mobile"],
            "Net Sales": [1200],
            "Order Number": ["ORD-1"],
            "Items Sold": [2],
            "Product Detail Views": [80],
            "Adds to Cart": [12],
            "Ecommerce Purchases": [2],
            "On Hand": [5],
            "Selling Price": [600],
            "COGS": [900],
            "Gross Margin": [300],
        })
        record = SimpleNamespace(dataframe=frame, role_mappings={})
        expected = {
            "brand": "Manufacturer",
            "product": "Item Name",
            "category": "Item Category",
            "revenue": "Net Sales",
            "order": "Order Number",
            "units": "Items Sold",
            "views": "Product Detail Views",
            "add_to_cart": "Adds to Cart",
            "transactions": "Ecommerce Purchases",
            "stock": "On Hand",
            "price": "Selling Price",
            "cost": "COGS",
            "margin": "Gross Margin",
        }
        numeric_roles = {"revenue", "units", "views", "add_to_cart", "transactions", "stock", "price", "cost", "margin"}
        for role, expected_column in expected.items():
            with self.subTest(role=role):
                resolved, candidates = resolve_column(record, role, numeric=role in numeric_roles)
                self.assertEqual((resolved, candidates), (expected_column, []))

    def test_four_commerce_quick_actions_use_real_rows_and_remain_read_only(self):
        workbook_id = self.upload()["workbook_id"]
        before = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()

        brand = self.execute(
            workbook_id,
            command="Revenue kolonuna göre en yüksek satış performansına sahip markayı bul; toplamını, payını ve ikinci markayla farkını göster",
            mode="commerce",
        )["result"]
        self.assertEqual(brand["metrics"]["leader"], "Apple")
        self.assertEqual(brand["metrics"]["runner_up"], "Samsung")
        self.assertAlmostEqual(brand["metrics"]["total"], 7000)
        self.assertAlmostEqual(brand["metrics"]["leader_share"], 4400 / 7000)
        self.assertAlmostEqual(brand["metrics"]["leader_gap"], 2300)
        self.assertEqual(brand["preview_rows"][0]["rank"], 1)
        self.assertIn("share", brand["preview_rows"][0])

        category = self.execute(
            workbook_id,
            command="Category bazında Revenue, sipariş ve ortalama sepet performansını karşılaştır; en güçlü ve en zayıf kategoriyi açıkla",
            mode="commerce",
        )["result"]
        self.assertEqual(category["metrics"]["strongest_category"], "Telefon")
        self.assertEqual(category["metrics"]["weakest_category"], "Accessories")
        self.assertEqual(category["metrics"]["total_orders"], 5)
        self.assertEqual(category["metrics"]["total_units"], 13)
        self.assertAlmostEqual(category["metrics"]["average_order_value"], 1400)
        telefon = next(row for row in category["preview_rows"] if row["Category"] == "Telefon")
        self.assertEqual(telefon["Revenue"], 3900)
        self.assertEqual(telefon["Order ID"], 2)
        self.assertEqual(telefon["Units"], 5)
        self.assertAlmostEqual(telefon["average_order_value"], 1950)

        funnel = self.execute(
            workbook_id,
            command="PDP View, Add to Cart ve Transaction kolonlarını kullanarak ürün bazında funnel dönüşümlerini hesapla ve en büyük kaybı bul",
            mode="commerce",
        )["result"]
        self.assertEqual(funnel["metrics"]["total_views"], 420)
        self.assertEqual(funnel["metrics"]["total_add_to_carts"], 90)
        self.assertEqual(funnel["metrics"]["total_transactions"], 13)
        self.assertEqual(funnel["metrics"]["largest_drop_stage"], "view_to_cart")
        self.assertEqual(funnel["metrics"]["largest_drop_count"], 330)
        self.assertEqual(funnel["metrics"]["weakest_product"], "MacBook")
        self.assertIn("view_to_purchase_rate", funnel["preview_rows"][0])
        self.assertIn("view_to_cart_drop_off", funnel["preview_rows"][0])

        stock = self.execute(
            workbook_id,
            command="Stock ve Revenue kolonlarını kullanarak yüksek talep gören kritik stoklu ürünleri önceliklendir",
            mode="commerce",
        )["result"]
        self.assertEqual(stock["metrics"]["total_revenue"], 7000)
        self.assertEqual(stock["metrics"]["total_stock"], 51)
        self.assertGreaterEqual(stock["metrics"]["at_risk_products"], 1)
        self.assertIn(stock["metrics"]["priority_product"], {"iPhone", "Watch"})
        self.assertEqual(stock["metrics"]["demand_metric"], "Units")
        self.assertTrue({"stock_risk", "priority_score", "revenue_per_stock"}.issubset(stock["preview_rows"][0]))

        after = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        self.assertFalse(after["is_modified"])
        self.assertEqual(after["preview"], before["preview"])

    def test_ambiguous_commerce_request_asks_then_accepts_explicit_column_answers(self):
        workbook_id = self.upload(ambiguous_revenue=True)["workbook_id"]
        first = self.execute(
            workbook_id,
            command="En performanslı markayı göster",
            mode="commerce",
        )
        self.assertEqual(first.get("status"), "needs_clarification", first)
        self.assertTrue(first.get("clarification_id"))
        self.assertTrue(first.get("question"))
        options = first.get("options") or []
        self.assertTrue(options)
        for option in options:
            self.assertTrue({"label", "value", "description"}.issubset(option), option)
        option_text = _all_scalar_text(options)
        self.assertIn("Revenue", option_text)

        second = self.execute(
            workbook_id,
            command="En performanslı markayı göster",
            mode="commerce",
            clarification_answers={
                "clarification_id": first["clarification_id"],
                "metric_column": "Net Revenue",
                "group_column": "Brand",
            },
        )
        self.assertEqual(second.get("status"), "success", second)
        self.assertIn("Net Revenue", second["result"]["columns_used"])
        self.assertIn("Brand", second["result"]["columns_used"])
        self.assertIn("Apple", _all_scalar_text(second["result"]))
        state = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        self.assertEqual(state["role_mappings"].get("metric_column"), "Net Revenue")

    def test_invalid_mode_and_unknown_structured_operation_fail_closed(self):
        workbook_id = self.upload()["workbook_id"]
        invalid_mode = self.client.post(
            "/api/excel-wizard/execute",
            json={"workbook_id": workbook_id, "command": "Revenue toplamı", "mode": "unsafe"},
        )
        self.assertEqual(invalid_mode.status_code, 422, invalid_mode.text)

        unknown = self.client.post(
            "/api/excel-wizard/execute",
            json={
                "workbook_id": workbook_id,
                "command": "Bilinmeyen işlem",
                "mode": "operations",
                "operation": {"id": "execute_arbitrary_python", "code": "raise SystemExit"},
            },
        )
        self.assertIn(unknown.status_code, {400, 422}, unknown.text)

    def test_failed_batch_rolls_back_every_prior_mutation(self):
        workbook_id = self.upload()["workbook_id"]
        response = self.client.post(
            "/api/excel-wizard/execute",
            json={
                "workbook_id": workbook_id,
                "command": "Önce kolon ekle, sonra geçersiz işlem yap",
                "mode": "operations",
                "operations": [
                    {"id": "upper", "source_column": "Brand", "target_column": "Must_Roll_Back"},
                    {"id": "data_validation", "column": "Status", "values": []},
                ],
            },
        )
        self.assertEqual(response.status_code, 400, response.text)
        state = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        self.assertFalse(state["is_modified"])
        self.assertNotIn("Must_Roll_Back", state["preview"]["columns"])

        workbook = load_workbook(
            BytesIO(self.client.get(f"/api/excel-wizard/download/{workbook_id}").content),
            data_only=False,
        )
        self.assertNotIn("Must_Roll_Back", [cell.value for cell in workbook["Commerce Data"][1]])

    def test_commerce_read_only_preflight_rejects_mutation_without_touching_workbook(self):
        workbook_id = self.upload()["workbook_id"]
        response = self.client.post(
            "/api/excel-wizard/execute",
            json={
                "workbook_id": workbook_id,
                "command": "Analiz modunda kolon eklemeye çalışma",
                "mode": "commerce",
                "operation": {"id": "upper", "source_column": "Brand", "target_column": "Must_Not_Exist"},
            },
        )
        self.assertEqual(response.status_code, 409, response.text)
        state = self.client.get("/api/excel-wizard/state", params={"workbook_id": workbook_id}).json()
        self.assertFalse(state["is_modified"])
        self.assertNotIn("Must_Not_Exist", state["preview"]["columns"])


class ExcelWizardV2UIMarkerTests(unittest.TestCase):
    def test_journey_uses_v2_endpoints_and_does_not_call_legacy_excel_routes(self):
        root = Path(main.__file__).resolve().parent
        journey = (root / "templates" / "journey.html").read_text(encoding="utf-8")
        script = (root / "static" / "excel-wizard.js").read_text(encoding="utf-8")
        combined = journey + "\n" + script
        for endpoint in (
            "/api/excel-wizard/upload",
            "/api/excel-wizard/state",
            "/api/excel-wizard/execute",
            "/api/excel-wizard/download/",
            "/api/excel-wizard/workbook/",
        ):
            self.assertIn(endpoint, combined)
        for legacy in ('fetch("/process-excel"', 'fetch("/upload-excel"', 'fetch("/active-excel-state"'):
            self.assertNotIn(legacy, combined)

    def test_ui_has_two_modes_clarification_controls_and_compact_result_region(self):
        root = Path(main.__file__).resolve().parent
        template = (root / "templates" / "excel_wizard.html").read_text(encoding="utf-8")
        journey = (root / "templates" / "journey.html").read_text(encoding="utf-8")
        style = (root / "static" / "excel-wizard.css").read_text(encoding="utf-8")
        script = (root / "static" / "excel-wizard.js").read_text(encoding="utf-8")
        combined = template + "\n" + journey + "\n" + style + "\n" + script

        # Stable semantic markers make the clarification flow accessible and
        # keep the result area compact enough that it does not dominate the UI.
        for marker in (
            'id="xwOperationsMode"',
            'id="xwCommerceMode"',
            'id="xwOperationsPanel"',
            'id="xwCommercePanel"',
            "xwClarificationCard",
            "data-xw-clarification-option",
            "xwClarificationAnswer",
            "xwClarificationSubmit",
            "xw-result-card",
            "xw-result-preview",
        ):
            self.assertIn(marker, combined)
        self.assertIn("needs_clarification", script)
        self.assertIn("clarification_answers", script)
        self.assertIn("data.role || data.field", script)
        self.assertIn("clarification_id: pending.id", script)
        self.assertIn("selected_value: answer", script)
        self.assertIn("hasResultPreview", script)
        self.assertIn("dataprovidoExcelWorkbookId", script)
        self.assertIn("sessionStorage", script)
        self.assertRegex(style, r"\.xw-result-(?:card|preview)[^\{]*\{")

        operations = re.findall(r"data-xw-operation\s*=\s*[\"']([^\"']+)", template)
        self.assertEqual(len(operations), 30, operations)
        self.assertEqual(len(set(operations)), 30, operations)
        self.assertIn("data-xw-capability-gallery", template)

    def test_ui_explains_ephemeral_processing_and_full_function_breadth(self):
        root = Path(main.__file__).resolve().parent
        template = (root / "templates" / "excel_wizard.html").read_text(encoding="utf-8")
        normalized = template.casefold()
        self.assertTrue(any(term in normalized for term in ("saklanmaz", "geçici", "temporary", "not stored")))
        for label in (
            "SUM",
            "IF",
            "SUMIFS",
            "XLOOKUP",
            "Pivot",
            "Koşullu Biçimlendirme",
            "Veri Doğrulama",
            "Grafik",
        ):
            self.assertIn(label.casefold(), normalized)

    def test_gallery_clicks_send_an_operation_hint_until_the_prompt_is_edited(self):
        root = Path(main.__file__).resolve().parent
        template = (root / "templates" / "excel_wizard.html").read_text(encoding="utf-8")
        journey = (root / "templates" / "journey.html").read_text(encoding="utf-8")
        script = (root / "static" / "excel-wizard.js").read_text(encoding="utf-8")

        operations = re.findall(r'data-xw-operation\s*=\s*["\']([^"\']+)', template)
        self.assertEqual(len(operations), 30, operations)
        self.assertIn("selectedOperation", script)
        self.assertIn("button.dataset.xwOperation", script)
        self.assertIn("getOperationHint()", script)
        self.assertIn("clearOperationHint()", script)
        self.assertIn("operation_hint:", journey)
        self.assertIn("window.ExcelWizardUI?.getOperationHint()", journey)


if __name__ == "__main__":
    unittest.main()
